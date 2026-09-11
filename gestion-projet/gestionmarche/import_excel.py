import argparse
import os
import re
import unicodedata
from datetime import datetime, date
from decimal import Decimal, InvalidOperation


from openpyxl import load_workbook, Workbook

DATE_HEADER_RE = re.compile(r'^\s*DATE\s*$', re.I)
TOTAL_RE = re.compile(r'^\s*TOTAL\s*$', re.I)
SURFACE_RE = re.compile(r'(\d+(?:[.,]\d+)?)\s*m(?:2|²)\b', re.I)
BLOCK_RE = re.compile(r'\b(EA|EB|EC|ED)\b', re.I)
DATE_STR_RE = re.compile(r'^\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*$')


def text(v):
    if v is None:
        return ''
    return str(v).strip()


def is_number(v):
    return isinstance(v, (int, float, Decimal)) and not isinstance(v, bool)


def amount(v):
    if is_number(v):
        return int(round(float(v)))
    s = text(v).replace('\xa0', '').replace(' ', '').replace(',', '.')
    if not s:
        return None
    try:
        d = Decimal(s)
        return int(d)
    except InvalidOperation:
        return None


def parse_date(v):
    if isinstance(v, datetime):
        return v.date(), False, ''
    if isinstance(v, date):
        return v, False, ''
    s = text(v)
    if not s:
        return None, True, 'Date absente'
    m = DATE_STR_RE.match(s)
    if not m:
        # One already-decided obvious typo: trailing text such as /cg.
        m2 = re.match(r'^\s*(\d{1,2})/(\d{1,2})/(\d{4})/[^0-9]+\s*$', s)
        if m2:
            d, mo, y = map(int, m2.groups())
            try:
                return date(y, mo, d), True, f"Date normalisée depuis '{s}'"
            except ValueError:
                return None, True, f"Date invalide '{s}'"
        return None, True, f"Date non reconnue '{s}'"
    d, mo, y = map(int, m.groups())
    try:
        return date(y, mo, d), True, ''
    except ValueError:
        return None, True, f"Date invalide '{s}'"


def normalize_name(s):
    s = unicodedata.normalize('NFKD', text(s))
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', s).strip().casefold()


def looks_like_cantine(s):
    s = text(s).upper()
    if not s:
        return False
    return bool(re.search(r'\bE[ABCD]\b|\bE[ABCD][A-Z0-9]|\b(?:SN|S/N)\b', s)) or bool(re.match(r'^\d+[\- ]', s))


def find_fiche_starts(rows):
    starts = []
    n = len(rows)
    for i in range(n):
        c = text(rows[i][2] if len(rows[i]) > 2 else None)
        if not c or DATE_HEADER_RE.match(c) or TOTAL_RE.match(c) or looks_like_cantine(c):
            continue
        # A fiche name is followed, within a few rows, by a DATE header.
        for j in range(i + 1, min(i + 6, n)):
            cj = text(rows[j][2] if len(rows[j]) > 2 else None)
            if DATE_HEADER_RE.match(cj):
                starts.append(i)
                break
    return starts


def parse_workbook(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    starts = find_fiche_starts(rows)
    starts_set = set(starts)
    fiches = []
    anonymous = []

    # Anonymous blocks are regions containing a DATE header but no recognized name start.
    covered = set()

    for pos, start in enumerate(starts):
        end = starts[pos + 1] if pos + 1 < len(starts) else len(rows)
        block = rows[start:end]
        name = text(rows[start][2] if len(rows[start]) > 2 else None)
        meta_parts = []
        superficie = None
        montant_total = None
        header_idx = None
        notes = []

        for rel, r in enumerate(block[1:], start=1):
            c = text(r[2] if len(r) > 2 else None)
            d = r[3] if len(r) > 3 else None
            e = r[4] if len(r) > 4 else None
            if DATE_HEADER_RE.match(c):
                header_idx = rel
                break
            if c:
                meta_parts.append(c)
            # surface can be in C or D
            for val in (c, text(d)):
                m = SURFACE_RE.search(val)
                if m:
                    try:
                        superficie = float(m.group(1).replace(',', '.'))
                    except ValueError:
                        pass
            # E is normally the total price. D can be price when C carries the surface.
            ea = amount(e)
            if ea is not None:
                montant_total = ea
            elif SURFACE_RE.search(c):
                da = amount(d)
                if da is not None:
                    montant_total = da

        if header_idx is None:
            notes.append('En-tête DATE introuvable')
            anonymous.append({'ligne': start + 1, 'nom': name, 'raison': 'Structure non reconnue'})
            continue

        # Keep the cantine reference while removing embedded surface notation only for display parsing.
        ref_parts = []
        for p in meta_parts:
            if p:
                ref_parts.append(p)
        numero = ' '.join(ref_parts).strip()
        # Remove standalone surface suffix only when it is clearly embedded; preserve original reference text otherwise.
        numero = re.sub(r'\s*/?\s*\d+(?:[.,]\d+)?\s*m(?:2|²)\b', '', numero, flags=re.I).strip(' /')

        prefixes = re.findall(r'\b(EA|EB|EC|ED)', numero.upper())
        prefixes = list(dict.fromkeys(prefixes))
        bloc = prefixes[0][1] if len(prefixes) == 1 else None
        if not bloc:
            notes.append('Bloc non déterminé automatiquement')

        payments = []
        control_total = None
        for rel in range(header_idx + 1, len(block)):
            r = block[rel]
            c = r[2] if len(r) > 2 else None
            d = r[3] if len(r) > 3 else None
            e = r[4] if len(r) > 4 else None
            cs = text(c)
            if TOTAL_RE.match(cs):
                control_total = amount(d)
                if control_total is None:
                    control_total = amount(e)
                break
            # Payment amount is in D for normal rows.
            amt = amount(d)
            if amt is None:
                # Blank/invalid amount: report only if there is a date-like value.
                if text(c):
                    dt, _, _ = parse_date(c)
                    if dt is not None or '/' in text(c):
                        notes.append(f"Versement sans montant à la ligne Excel {start + rel + 1}")
                continue
            dt, hour_unknown, date_note = parse_date(c)
            pnote = date_note
            payments.append({
                'excel_row': start + rel + 1,
                'date': dt,
                'hour_unknown': True,
                'amount': amt,
                'note': pnote,
            })
            if date_note:
                notes.append(f"Ligne {start + rel + 1}: {date_note}")

        detail_sum = sum(p['amount'] for p in payments)
        if control_total is not None and control_total != detail_sum:
            notes.append(f"TOTAL Excel {control_total:,} ≠ somme détaillée {detail_sum:,}")

        if not name:
            anonymous.append({'ligne': start + 1, 'nom': '', 'raison': 'Nom absent'})
            continue
        if not numero:
            notes.append('Numéro/référence de cantine absent')
        fiches.append({
            'excel_row': start + 1,
            'nom': name,
            'nom_normalise': normalize_name(name),
            'numero_cantine': numero,
            'superficie': superficie,
            'montant_a_payer': montant_total,
            'bloc_suggere': bloc,
            'notes': ' | '.join(notes),
            'control_total': control_total,
            'detail_sum': detail_sum,
            'payments': payments,
        })
        covered.update(range(start, end))

    # Scan for unassigned DATE blocks near the end / anywhere not covered by a fiche.
    i = 0
    while i < len(rows):
        if i in covered:
            i += 1
            continue
        c = text(rows[i][2] if len(rows[i]) > 2 else None)
        if DATE_HEADER_RE.match(c):
            # Capture until TOTAL or next blank run.
            j = i + 1
            payment_count = 0
            while j < len(rows):
                cj = text(rows[j][2] if len(rows[j]) > 2 else None)
                if TOTAL_RE.match(cj):
                    break
                if amount(rows[j][3] if len(rows[j]) > 3 else None) is not None:
                    payment_count += 1
                if not cj and payment_count == 0:
                    break
                j += 1
            anonymous.append({'ligne': i + 1, 'nom': '', 'raison': f'Bloc sans nom (versements détectés: {payment_count})'})
            i = max(i + 1, j + 1)
        else:
            i += 1

    # Duplicate-name and duplicate-cantine diagnostics.
    name_counts = {}
    ref_counts = {}
    for f in fiches:
        name_counts.setdefault(f['nom_normalise'], []).append(f['excel_row'])
        if f['numero_cantine']:
            ref_counts.setdefault(f['numero_cantine'].casefold(), []).append(f['excel_row'])
    for f in fiches:
        if len(name_counts[f['nom_normalise']]) > 1:
            f['notes'] = (f['notes'] + ' | ' if f['notes'] else '') + 'Nom présent sur plusieurs fiches : fusion à contrôler'
        if f['numero_cantine'] and len(ref_counts[f['numero_cantine'].casefold()]) > 1:
            f['notes'] = (f['notes'] + ' | ' if f['notes'] else '') + 'Référence de cantine répétée : contrôle requis'

    return fiches, anonymous


def write_report(out, fiches, anonymous):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Fiches'
    headers = ['Ligne Excel','Client','Numéro de cantine','Superficie','Montant à payer','Bloc suggéré','Nb versements','Somme versements','TOTAL Excel contrôle','Notes']
    ws.append(headers)
    for f in fiches:
        ws.append([
            f['excel_row'], f['nom'], f['numero_cantine'], f['superficie'], f['montant_a_payer'], f['bloc_suggere'],
            len(f['payments']), f['detail_sum'], f['control_total'], f['notes']
        ])

    wp = wb.create_sheet('Versements')
    wp.append(['Ligne Excel','Client','Numéro de cantine','Date','Heure','Montant','Note'])
    for f in fiches:
        for p in f['payments']:
            wp.append([p['excel_row'], f['nom'], f['numero_cantine'], p['date'], 'Heure inconnue', p['amount'], p['note']])

    wa = wb.create_sheet('Anonymes')
    wa.append(['Ligne Excel','Client','Raison'])
    for a in anonymous:
        wa.append([a['ligne'], a['nom'], a['raison']])

    wd = wb.create_sheet('Décisions appliquées')
    wd.append(['Règle','Traitement'])
    decisions = [
        ('TOTAL Excel','Ignoré comme versement ; utilisé uniquement comme valeur de contrôle.'),
        ('Incohérence Bathie Niasse','Les versements détaillés sont conservés ; aucun versement artificiel n’est créé.'),
        ('Incohérence Mafall Fall','Les versements détaillés sont conservés ; aucun versement artificiel n’est créé.'),
        ('Date absente','Date inconnue + heure inconnue.'),
        ('Date connue sans heure','Date conservée + heure inconnue.'),
        ('Nouveau versement dans l’application','Date et heure automatiques au moment de l’enregistrement.'),
        ('Fiche sans nom','Non importée automatiquement.'),
        ('Nom présent sur plusieurs fiches','Fusion/association à contrôler ; aucune fusion ambiguë n’est imposée dans l’aperçu.'),
    ]
    for row in decisions:
        wd.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = 'A2'
        for col in sheet.columns:
            maxlen = max(len(str(c.value or '')) for c in col)
            sheet.column_dimensions[col[0].column_letter].width = min(max(maxlen + 2, 12), 60)

    wb.save(out)


def main():
    ap = argparse.ArgumentParser(description='Import sécurisé des fiches de versement Excel dans Gestion du Marché.')
    ap.add_argument('excel', nargs='?', default='FICHE VERSEEMNT(1).xlsx', help='Fichier Excel source')
    ap.add_argument('--preview', action='store_true', help='Aperçu uniquement (mode par défaut)')
    ap.add_argument('--import-mode', action='store_true', help="Effectue l'import réel dans la base Django")
    ap.add_argument('--report', default='rapport_import_preview.xlsx', help='Rapport Excel de prévisualisation')
    args = ap.parse_args()

    if not os.path.exists(args.excel):
        raise SystemExit(f"Fichier introuvable : {args.excel}")

    fiches, anonymous = parse_workbook(args.excel)
    write_report(args.report, fiches, anonymous)

    payments = sum(len(f['payments']) for f in fiches)
    flagged = sum(bool(f['notes']) for f in fiches)

    if args.import_mode:
        import django
        os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'gestionmarche.settings')
        django.setup()
        from django.db import transaction
        from gestion.models import Client, Cantine, Paiement

        # Cas explicitement laissé à contrôler : même nom, deux situations distinctes.
        AMBIGUOUS_NAMES = {'ousseynou sall'}
        created_clients = created_cantines = created_payments = 0
        reused_clients = reused_cantines = 0
        skipped = []

        def find_client(name):
            key = normalize_name(name)
            matches = [c for c in Client.objects.all() if normalize_name(c.nom) == key]
            if len(matches) == 1:
                return matches[0], False
            if len(matches) > 1:
                return None, False
            c = Client.objects.create(nom=name, telephone='')
            return c, True

        with transaction.atomic():
            for f in fiches:
                name_key = f['nom_normalise']
                if name_key in AMBIGUOUS_NAMES:
                    skipped.append(f"Ligne {f['excel_row']} — {f['nom']} : cas ambigu laissé à contrôler")
                    continue
                client, was_created = find_client(f['nom'])
                if client is None:
                    skipped.append(f"Ligne {f['excel_row']} — {f['nom']} : plusieurs clients existants portent ce nom")
                    continue
                if was_created:
                    created_clients += 1
                else:
                    reused_clients += 1

                ref = f['numero_cantine']
                if not ref:
                    skipped.append(f"Ligne {f['excel_row']} — {f['nom']} : référence de cantine absente")
                    continue

                existing = Cantine.objects.filter(numero_cantine=ref).first()
                if existing:
                    if existing.client_id not in (None, client.id):
                        skipped.append(f"Ligne {f['excel_row']} — {f['nom']} / {ref} : référence déjà liée à un autre client")
                        continue
                    cantine = existing
                    if existing.client_id is None:
                        existing.client = client
                        existing.save(update_fields=['client'])
                    reused_cantines += 1
                else:
                    bloc = f['bloc_suggere'] or 'B'
                    note = f['notes']
                    if not f['bloc_suggere']:
                        note = (note + ' | ' if note else '') + 'Bloc non déterminé : B utilisé provisoirement'
                    superficie = f['superficie']
                    # Le modèle actuel attend une chaîne pour superficie.
                    superficie_value = '' if superficie is None else str(superficie).rstrip('0').rstrip('.') if isinstance(superficie, float) else str(superficie)
                    cantine = Cantine.objects.create(
                        client=client,
                        numero_cantine=ref,
                        superficie=superficie_value,
                        montant_a_payer=f['montant_a_payer'] or 0,
                        bloc=bloc,
                        notes=note,
                    )
                    created_cantines += 1

                for p in f['payments']:
                    # Évite un doublon évident si le script est relancé.
                    dt = p['date']
                    from django.utils import timezone
                    if dt is None:
                        candidate = Paiement.objects.filter(cantine=cantine, montant_paye=p['amount'], date_paiement__isnull=True).first()
                    else:
                        candidate = Paiement.objects.filter(cantine=cantine, montant_paye=p['amount'], date_paiement__date=dt).first()
                    if candidate:
                        continue
                    payment = Paiement(cantine=cantine, montant_paye=p['amount'], heure_inconnue=True)
                    if dt is not None:
                        payment.date_paiement = timezone.make_aware(datetime.combine(dt, datetime.min.time()))
                    payment.save()
                    created_payments += 1

        print('IMPORT RÉEL TERMINÉ')
        print(f'Clients créés : {created_clients}')
        print(f'Clients réutilisés : {reused_clients}')
        print(f'Cantines créées : {created_cantines}')
        print(f'Cantines réutilisées : {reused_cantines}')
        print(f'Versements créés : {created_payments}')
        print(f'Cas non importés / à contrôler : {len(skipped)}')
        for item in skipped:
            print(' -', item)
        print(f'Rapport : {args.report}')
    else:
        print('MODE APERÇU — AUCUNE MODIFICATION DE LA BASE')
        print(f'Fiches détectées : {len(fiches)}')
        print(f'Versements détectés : {payments}')
        print(f'Fiches nécessitant un contrôle : {flagged}')
        print(f'Blocs sans nom : {len(anonymous)}')
        print(f'Rapport créé : {args.report}')


if __name__ == '__main__':
    main()
